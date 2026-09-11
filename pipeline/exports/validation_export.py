# -*- coding: utf-8 -*-
"""ไฟล์รายการปัญหาที่ต้องไล่แก้ — .xlsx และ .csv

ค่าเริ่มต้นเอาเฉพาะ "ข้อมูลปัจจุบันที่ยังไม่ได้ดู" (scope=current, status=NEW)
คือรายการที่ยังตามกลับไปแก้ต้นทางได้จริง

ในไฟล์ไม่มีชื่อ เบอร์โทร หรือที่อยู่ลูกค้า — มีแค่เลขที่ออเดอร์กับต้นทาง
จึงเอาไปส่งต่อในทีมได้ (แต่ data/output/ ถูก .gitignore กันไว้อยู่แล้ว)
"""
from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("rule", "กฎที่ตรวจ", 26),
    ("entity_id", "เลขที่ออเดอร์", 24),
    ("entity_source", "ต้นทาง", 22),
    ("entity_date", "วันที่ออเดอร์", 13),
    ("scope", "ประเภทข้อมูล", 12),
    ("status", "สถานะ", 10),
    ("severity", "ระดับ", 9),
    ("detail", "ปัญหาที่เจอ", 62),
    ("seen_count", "เจอมาแล้วกี่รอบ", 15),
    ("first_seen_at", "เจอครั้งแรก", 20),
    ("last_seen_at", "เจอล่าสุด", 20),
    ("status_note", "บันทึกของผู้ตรวจ", 30),
]

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def _rows(con, scope: str = "current", status: str = "NEW") -> list[dict]:
    sql = "SELECT * FROM validation_issue WHERE 1=1"
    params: list = []
    if scope:
        sql += " AND COALESCE(scope,'current')=?"
        params.append(scope)
    if status:
        sql += " AND status=?"
        params.append(status)
    sql += " ORDER BY rule, entity_date DESC, entity_id"
    return [dict(r) for r in con.execute(sql, params)]


def write(con, out_dir: Path, date_tag: str, scope: str = "current",
          status: str = "NEW") -> dict:
    """เขียนทั้ง .xlsx และ .csv คืนสรุปว่าเขียนอะไรไปบ้าง"""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = _rows(con, scope, status)

    xlsx_path = out_dir / f"validation_issues_{date_tag}.xlsx"
    csv_path = out_dir / f"validation_issues_{date_tag}.csv"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "issues"
    ws.append([label for _, label, _ in COLUMNS])
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([r.get(key) for key, _, _ in COLUMNS])
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"
    wb.save(xlsx_path)
    wb.close()

    with open(csv_path, "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow([key for key, _, _ in COLUMNS])
        for r in rows:
            w.writerow([r.get(key) for key, _, _ in COLUMNS])

    by_rule: dict[str, int] = {}
    for r in rows:
        by_rule[r["rule"]] = by_rule.get(r["rule"], 0) + 1
    return {"rows": len(rows), "xlsx": xlsx_path, "csv": csv_path, "by_rule": by_rule}
