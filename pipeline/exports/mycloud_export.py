# -*- coding: utf-8 -*-
"""สร้างไฟล์ .xlsx สำหรับอัปโหลดเข้า MyCloud

พอร์ตจาก index.html: buildMcRows() + mcFillTemplate()

ของเดิมเขียน ZIP reader, ตาราง CRC-32, deflate และตัวแก้ XML ของ xlsx เองราว 300 บรรทัด
เพราะเบราว์เซอร์เปิดไฟล์ template ไม่ได้ด้วยวิธีอื่น ที่นี่ใช้ openpyxl เปิด template ตรง ๆ

กติกาที่ยกมาทั้งหมด
  * ราคา = ราคาต่อหน่วยหลังหักส่วนลด รวม VAT ไม่ใช่ราคา catalog
  * ออเดอร์ที่ไม่มีคำว่า COD หรือ โอน ใน Lead Intro -> ข้ามทั้งใบ ไม่เดา
  * Delivery mode = (COD หรือ โอน) x (กรุงเทพปริมณฑล หรือ ต่างจังหวัด)
  * เฉพาะบรรทัดแรกของออเดอร์ที่มีชื่อ/เบอร์/ที่อยู่/รหัสไปรษณีย์/ขนส่ง/สถานะ
    บรรทัดถัดไปใส่แค่ SKU / Price / QTY
  * Fulfill ก็ต่อเมื่อสินค้าอนุมัติครบทุกตัว มีตัวใดยังจอง -> ทั้งใบเป็น OnHold
  * Olive Oil 1 ขวด และเป็นสินค้าตัวเดียวในออเดอร์ -> ราคา 300 (รวมค่าส่ง)
  * SKU ซ้ำในออเดอร์เดียวกันต้องยุบเหลือแถวเดียว (MyCloud ปฏิเสธไฟล์ที่ SKU ซ้ำ)
"""
from __future__ import annotations

import re
import shutil
import sqlite3
from pathlib import Path

import openpyxl

from .. import normalize as nz
from .. import util
from ..config import Config, ROOT

TEMPLATE_PATH = ROOT / "config" / "templates" / "mycloud_template.xlsx"
SHEET_NAME = "Template"
START_ROW = 3                    # แถว 1 = รหัส IM, แถว 2 = หัวตาราง

# ฟิลด์ -> คอลัมน์ (index.html: MC_FIELD_COL)
FIELD_COL = {
    "name": "A",             # IM01 ชื่อผู้รับ
    "phone": "C",            # IM03 เบอร์โทรศัพท์
    "address": "E",          # IM05 ที่อยู่ผู้รับ
    "postcode": "G",         # IM07 รหัสไปรษณีย์
    "sku": "I",              # IM09 Shop SKU
    "price": "K",            # IM11 ราคา (ตัวเลข)
    "qty": "L",              # IM12 จำนวน (ตัวเลข)
    "delivery_mode": "N",    # IM14 ขนส่ง
    "status": "R",           # IM18 สถานะ
    "shop_order_no": "U",    # IM21 Shop Order Number (= order_id ไว้ตามรอย)
    "sales_channel": "V",    # IM22 ช่องทางขาย
    "sales_person": "W",     # IM23 ผู้ขาย
    "notes": "X",            # IM24 หมายเหตุภายในร้าน
}
NUMERIC_FIELDS = {"price", "qty"}


def map_sales_channel(cfg: Config, lead_intro: str, owner: str) -> tuple[str, str]:
    """ช่องทางขายสำหรับไฟล์ MC (index.html: mapSalesChannel) — คืน (รหัส, หมายเหตุ)

    อันดับ 1 ดูจาก Lead Intro | อันดับ 2 ดูจากผู้ดูแล
    """
    ch_codes = cfg.ch
    intro = str(lead_intro or "")
    ch = nz.parse_channel(intro)

    if re.search(r"instagram", ch, re.I):
        return ch_codes["ig"], ""
    if "ไขมัน" in ch:
        return ch_codes["fb_kaimun"], ""
    if "ผิว" in ch:
        return ch_codes["fb_piw"], ""
    if "นอน" in ch:
        return ch_codes["fb_non"], ""
    if re.search(r"facebook", ch, re.I):
        return ch_codes["fb_hylme"], ""
    # marketplace ไม่ได้ส่งของผ่าน MC -> เว้นช่องทางขายว่าง และไม่ถือเป็นข้อมูลไม่ครบ
    if cfg.re_marketplace.search(intro) or cfg.re_marketplace.search(ch):
        return "", ""
    if re.search(r"telesales|เทเลเซล|โทรขาย", intro, re.I) or re.search(r"telesales", ch, re.I):
        return ch_codes["telesales"], ""
    if re.search(r"line|ไลน์", intro, re.I):
        return (ch_codes["line_privilege"] if re.search(r"privilege", intro, re.I)
                else ch_codes["line_official"]), ""
    if re.search(r"(^|[^A-Za-z])CRM([^A-Za-z]|$)", intro, re.I):
        return ch_codes["line_privilege"], ""

    person = util.norm(owner)
    if person in cfg.telesales_people:
        return ch_codes["telesales"], ""
    return "", "ระบุช่องทางขายจาก Lead Intro/ผู้ดูแลไม่ได้ กรุณากรอกช่องทางขายเอง"


def build_rows(con: sqlite3.Connection, cfg: Config, date_from: str, date_to: str) -> dict:
    """สร้างแถวสำหรับไฟล์ MC จากออเดอร์ในฐานข้อมูล คืน {rows, skipped, flagged}"""
    orders = con.execute(
        "SELECT * FROM orders WHERE order_date BETWEEN ? AND ? "
        "AND COALESCE(is_cancelled,0) = 0 ORDER BY order_ts, order_id",
        (date_from, date_to),
    ).fetchall()

    rows: list[dict] = []
    skipped: list[dict] = []
    flagged: list[dict] = []

    for o in orders:
        order_id = o["order_id"]
        source = str(o["source"] or "")
        # ออเดอร์ที่อ่านกลับมาจาก MyCloud/marketplace ไม่ต้อง export กลับเข้า MC
        # (ไม่มี SKU และชื่อ/เบอร์/ที่อยู่ถูกปิดบัง)
        if re.search(r"MyCloud|Marketplace", source, re.I):
            continue

        # วิธีชำระเงิน: อ่านจาก Lead Intro ก่อนตามกติกาเดิม
        # ออเดอร์เก่าที่โหลดเข้ามาก่อนมี pipeline ไม่มี Lead Intro แต่มีคอลัมน์
        # payment_method ที่คำนวณไว้แล้ว จึงใช้เป็นตัวสำรอง
        pm = nz.payment_method(o["lead_intro"])
        if pm == nz.PAYMENT_UNKNOWN:
            stored = util.norm(o["payment_method"])
            if stored in (nz.PAYMENT_COD, nz.PAYMENT_TRANSFER):
                pm = stored
        if pm == nz.PAYMENT_UNKNOWN:
            reason = ('ไม่มีคำว่า "COD" หรือ "โอน" ใน Lead Intro — ไม่ทราบวิธีชำระเงิน'
                      if o["lead_intro"] else
                      'ออเดอร์เก่าที่ยังไม่ผ่าน pipeline (ไม่มี Lead Intro และไม่มีวิธีชำระเงินบันทึกไว้)')
            skipped.append({"order_id": order_id, "name": o["name_raw"], "reason": reason})
            continue

        items = con.execute(
            "SELECT * FROM order_items WHERE order_id=? AND sku IS NOT NULL AND sku <> '' "
            "ORDER BY line_seq", (order_id,)
        ).fetchall()
        if not items:
            skipped.append({
                "order_id": order_id, "name": o["name_raw"],
                "reason": "ไม่มีรายการสินค้าที่มี Product Code (SKU) ในไฟล์",
            })
            continue

        name = nz.build_recipient_name("", "", o["recipient_name"] or o["name_raw"])
        phone = o["phone_clean"] or ""
        address = o["address"] or ""
        province, prov_source = nz.resolve_province(o["province"], address)
        postcode = o["postcode"] or (util.find_zip_in_text(address) or "")

        notes: list[str] = []
        if not name:
            notes.append("ไม่มีชื่อผู้รับ")
        elif util.is_masked(name):
            notes.append("ชื่อผู้รับถูกปิดบัง (มี *) — ต้องแก้ก่อนส่ง")
        if not phone:
            notes.append("ไม่มีเบอร์โทรที่ใช้ได้")
        if not address:
            notes.append("ไม่มีที่อยู่")
        if not postcode:
            notes.append("ไม่มีรหัสไปรษณีย์")
        if prov_source == "unknown":
            notes.append("หาจังหวัดไม่พบ ระบบใช้ค่าเริ่มต้น = ต่างจังหวัด (EMS) กรุณาตรวจสอบ")

        is_cod = (pm == nz.PAYMENT_COD)
        mode = nz.delivery_mode(cfg, is_cod, province, prov_source)

        # ---- สถานะ: อนุมัติครบทุกตัวเท่านั้นจึงเป็น Fulfill ----
        status = (cfg.statuses["status_fulfill"]
                  if all(nz.is_sku_approved(cfg, it["sku"]) for it in items)
                  else cfg.statuses["status_onhold"])

        extra_notes = []
        preorder = [it for it in items if nz.is_preorder_item(cfg, it["sku"], it["product_raw"])]
        if preorder:
            names = list(dict.fromkeys(nz.short_product_name(cfg, it["product_raw"]) for it in preorder))
            extra_notes.append("Pre-order: " + ", ".join(names))

        channel_code, chan_note = map_sales_channel(cfg, o["lead_intro"], o["owner"])
        if chan_note:
            notes.append(chan_note)
        person = (cfg.mappings["telesales_mc_person"]
                  if channel_code == cfg.ch["telesales"]
                  else nz.map_sales_person(cfg, o["owner"]))
        all_notes = notes + extra_notes

        # ---- ราคา: ปัดเป็นจำนวนเต็มโดยผลรวมยังเท่ายอดที่เก็บลูกค้าจริง ----
        line_qty = [util.num(it["qty"]) for it in items]
        line_raw = []
        for i, it in enumerate(items):
            if len(items) == 1 and nz.is_olive_oil(cfg, it["sku"], it["product_raw"]) and line_qty[i] == 1:
                line_raw.append(float(cfg.products["olive_one_bottle_price"]))
            else:
                line_raw.append(util.with_vat(it["amount_ex_vat"], cfg.vat_rate))
        line_totals = util.allocate_whole_baht(line_raw)

        # ---- ยุบ SKU ที่ซ้ำในออเดอร์เดียวกันให้เหลือแถวเดียว ----
        # MyCloud ปฏิเสธไฟล์ที่มี Shop SKU ซ้ำภายในคำสั่งซื้อเดียวกัน
        by_sku: dict[str, dict] = {}
        for i, it in enumerate(items):
            sku = it["sku"] or ""
            g = by_sku.get(sku) or {"total": 0.0, "qty": 0.0, "order": len(by_sku)}
            g["total"] += line_totals[i]
            g["qty"] += line_qty[i]
            by_sku[sku] = g

        emitted = 0
        for sku, g in sorted(by_sku.items(), key=lambda kv: kv[1]["order"]):
            first = (emitted == 0)
            emitted += 1
            price = round((g["total"] / g["qty"]) * 100) / 100 if g["qty"] > 0 else g["total"]
            rows.append({
                "name": name if first else "",
                "phone": phone if first else "",
                "address": address if first else "",
                "postcode": postcode if first else "",
                "sku": sku,
                "price": price,
                "qty": g["qty"],
                "delivery_mode": mode if first else "",
                "status": status if first else "",
                "shop_order_no": order_id if first else "",
                "sales_channel": channel_code if first else "",
                "sales_person": person if first else "",
                "notes": "; ".join(all_notes) if first else "",
            })
        if notes:
            flagged.append({"order_id": order_id, "name": name, "notes": "; ".join(all_notes)})

    return {"rows": rows, "skipped": skipped, "flagged": flagged}


def write_xlsx(rows: list[dict], out_path: Path) -> Path:
    """เติมข้อมูลลง template แล้วบันทึกเป็นไฟล์ใหม่

    เปิด template ด้วย openpyxl แล้วเขียนทับเฉพาะเซลล์ที่ต้องใส่
    หัวตาราง สไตล์ และ dropdown ของ template ยังอยู่ครบ
    """
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"ไม่พบไฟล์ template ที่ {TEMPLATE_PATH}\n"
            f"ไฟล์นี้ถูกดึงออกมาจาก index.html แล้ววางไว้ที่ config/templates/"
        )
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE_PATH, out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
    for i, r in enumerate(rows):
        row_num = START_ROW + i
        for field, col in FIELD_COL.items():
            v = r.get(field)
            if v is None or v == "":
                continue
            cell = ws[f"{col}{row_num}"]
            cell.value = float(v) if field in NUMERIC_FIELDS else str(v)
    wb.save(out_path)
    wb.close()
    return out_path
