# -*- coding: utf-8 -*-
"""ตัวอ่านไฟล์ Excel แต่ละชนิด

แต่ละ loader รับ path ของไฟล์ แล้วคืน LoadResult ที่มีโครงสร้างกลางเหมือนกันหมด
เพื่อให้ขั้นตอนถัดไป (normalize -> db) ไม่ต้องรู้ว่าไฟล์มาจากระบบไหน
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl


@dataclass
class LoadResult:
    """ผลการอ่านไฟล์หนึ่งไฟล์"""
    kind: str                                   # pancake | mycloud | marketplace | cancel | unknown
    file_name: str
    sheet: str = ""
    orders: list[dict] = field(default_factory=list)
    items: list[dict] = field(default_factory=list)      # มี order_id ชี้กลับไปที่ orders
    events: list[dict] = field(default_factory=list)     # เหตุการณ์การจัดส่ง
    cancels: list[dict] = field(default_factory=list)    # ทะเบียนออเดอร์ยกเลิก
    sources: list[dict] = field(default_factory=list)    # (system, external_id, order_id)
    page_stat: dict | None = None                       # สถิติหน้า Pancake (ตัวหารของ Conversion)
    sales_report: list[dict] | None = None              # เป้ายอดขายรายวัน x ช่องทาง จากรายงาน Excel
    audit: list[dict] = field(default_factory=list)     # ที่มาของแต่ละออเดอร์ (ไฟล์/ชีต/แถว/mapping)
    replace_scope: dict | None = None                   # {"source":..., "dates":[...]} ลบของเดิมก่อนโหลด
    stats: dict = field(default_factory=dict)
    rows_read: int = 0

    @property
    def date_range(self) -> tuple[str | None, str | None]:
        dates = sorted(o["order_date"] for o in self.orders if o.get("order_date"))
        if not dates and self.page_stat:
            dates = list(self.page_stat.get("date_list") or [])
        if not dates and self.cancels:
            dates = sorted(c["cancel_date"] for c in self.cancels if c.get("cancel_date"))
        return (dates[0], dates[-1]) if dates else (None, None)


# ---------------------------------------------------------------------------
# เปิดไฟล์ค้างไว้ทีละหนึ่งไฟล์
# ---------------------------------------------------------------------------
# openpyxl ใช้เวลาส่วนใหญ่ไปกับ load_workbook (ไฟล์จริง 14 MB = 13 วินาที) ส่วนการ
# วนอ่านชีตหลังเปิดแล้วเร็วมาก (0.02 วินาที) และอ่านซ้ำได้ไม่จำกัด
# ของเดิมทุกฟังก์ชันเปิด-ปิดไฟล์เองทุกครั้ง ตัวตรวจชนิดไฟล์ที่ไล่ดูทีละชีตจึงเปิดไฟล์
# ซ้ำ 31 รอบ = 7 นาทีต่อไฟล์เดียว แคชนี้ทำให้เหลือ 13 วินาที
# แคชทีละไฟล์พอ เพราะ pipeline ทำงานทีละไฟล์อยู่แล้ว และคีย์รวมเวลาแก้ไข+ขนาด
# ถ้าไฟล์ถูกแก้ระหว่างรัน จะเปิดใหม่ให้เอง
_WB_CACHE: tuple | None = None


def open_workbook(path: Path):
    """เปิดไฟล์แบบใช้ซ้ำได้ — อย่าเรียก .close() บนผลลัพธ์ ใช้ close_workbook() แทน"""
    global _WB_CACHE
    path = Path(path)
    stat = path.stat()
    key = (str(path.resolve()), stat.st_mtime_ns, stat.st_size)
    if _WB_CACHE is not None and _WB_CACHE[0] == key:
        return _WB_CACHE[1]
    close_workbook()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    _WB_CACHE = (key, wb)
    return wb


def close_workbook() -> None:
    """คืนหน่วยความจำของไฟล์ที่เปิดค้างไว้ — เรียกเมื่ออ่านไฟล์นั้นเสร็จแล้ว"""
    global _WB_CACHE
    if _WB_CACHE is not None:
        try:
            _WB_CACHE[1].close()
        except Exception:
            pass
        _WB_CACHE = None


def read_sheet_dicts(path: Path, sheet_name: str) -> list[dict]:
    """อ่านชีตเป็น list ของ dict โดยใช้แถวแรกเป็นหัวคอลัมน์

    เทียบเท่า XLSX.utils.sheet_to_json(ws, {defval: null}) ใน index.html
    ช่องว่างคืนค่า None เหมือนกัน
    """
    ws = open_workbook(path)[sheet_name]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    cols = [(str(h).strip() if h is not None else "") for h in header]
    out = []
    for r in rows:
        if r is None or all(v is None for v in r):
            continue
        out.append({cols[i]: r[i] for i in range(min(len(cols), len(r))) if cols[i]})
    return out


def read_sheet_matrix(path: Path, sheet_name: str, max_rows: int | None = None) -> list[list]:
    """อ่านชีตเป็นตารางดิบ (ใช้ตอนหาแถวหัวตารางที่ไม่ได้อยู่บรรทัดแรก)"""
    ws = open_workbook(path)[sheet_name]
    out = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if max_rows is not None and i >= max_rows:
            break
        out.append(list(r))
    return out


def sheet_names(path: Path) -> list[str]:
    return list(open_workbook(path).sheetnames)


def sheet_names_fast(path: Path) -> list[str]:
    """ชื่อชีตโดยไม่ต้องเปิดไฟล์ด้วย openpyxl — อ่าน xl/workbook.xml จาก zip ตรง ๆ

    ทำไมต้องมี: openpyxl ต้องแกะ sharedStrings ทั้งก้อนก่อนถึงจะบอกชื่อชีตได้
    ไฟล์จริงขนาด 15 MB ใช้เวลาราว 10 วินาทีต่อการเปิดหนึ่งครั้ง ถ้าตัวตรวจชนิดไฟล์
    เปิดกันคนละที แค่ขั้นตอน "ดูว่าไฟล์นี้เป็นชนิดไหน" ก็กินเวลาเป็นนาที
    ฟังก์ชันนี้ใช้เวลาระดับมิลลิวินาที ใช้เป็นด่านแรกของ detect() เพื่อคัดไฟล์ที่ไม่ใช่ออกก่อน
    """
    import re as _re
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("xl/workbook.xml").decode("utf-8", "replace")
    except Exception:
        return sheet_names(path)          # .xls เก่า หรือไฟล์ที่ zip อ่านไม่ได้
    out = []
    for m in _re.finditer(r"<(?:\w+:)?sheet\b[^>]*?\bname\s*=\s*\"([^\"]*)\"", xml):
        s = m.group(1)
        for a, b in (("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'),
                     ("&apos;", "'"), ("&amp;", "&")):
            s = s.replace(a, b)
        out.append(s)
    return out


def field_of(row: dict, names: list[str]):
    """อ่านค่าจากแถวโดยชื่อคอลัมน์ รองรับหลายชื่อ ช่องว่างไม่สำคัญ

    index.html: pkField — คืน None เมื่อไม่มีคอลัมน์นั้นหรือค่าว่าง
    (การแยก 'ไม่มีคอลัมน์' ออกจาก 'มีคอลัมน์แต่เป็น 0' สำคัญมากกับออเดอร์แจกฟรี)
    """
    import re as _re
    for n in names:
        for key, v in row.items():
            if _re.sub(r"\s+", " ", str(key)).strip() == n:
                if v is not None and str(v).strip() != "":
                    return v
    return None


def num_of(row: dict, names: list[str]) -> float:
    """index.html: pkNum"""
    from .. import util
    v = field_of(row, names)
    return 0.0 if v is None else util.num(v)
